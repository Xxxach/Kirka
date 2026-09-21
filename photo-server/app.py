import io
import os
import time
import textwrap
import urllib.request
from collections import Counter

import cv2
import numpy as np
import pdfplumber
import rinoh_typeface_dejavusans as _dejavu_pkg
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from PIL import Image, ImageFilter
from rembg import remove
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas

FRONTEND_ORIGIN = os.environ.get("FRONTEND_ORIGIN", "http://localhost:5173")

app = Flask(__name__)
CORS(app, origins=[FRONTEND_ORIGIN])
app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024  # 15 МБ на файл

QUALITY_MAP = {"low": 85, "medium": 60, "high": 35}

CONVERT_FORMATS = {
    "png": ("PNG", "image/png"),
    "jpg": ("JPEG", "image/jpeg"),
    "webp": ("WEBP", "image/webp"),
    "ico": ("ICO", "image/x-icon"),
}

DAILY_LIMIT = 70  # бесплатных обработок фото на IP в сутки
DAY_SECONDS = 24 * 60 * 60
MAX_INPUT_SIDE = 2000  # защита слабого сервера от гигантских исходников

usage = {}

# ---------- Документы: PDF <-> Word (только текст, без таблиц/картинок) ----------
# Шрифт с поддержкой кириллицы для генерации PDF — иначе reportlab рисует
# только латиницу.
_FONT_DIR = os.path.dirname(_dejavu_pkg.__file__)
pdfmetrics.registerFont(TTFont("DejaVuSans", os.path.join(_FONT_DIR, "DejaVuSans.ttf")))


def extract_lines_with_bbox(page):
    """Группирует слова PDF-страницы в строки с их горизонтальными
    границами — нужно, чтобы понять, была ли строка по центру."""
    words = page.extract_words()
    words.sort(key=lambda w: (w["top"], w["x0"]))
    lines = []
    current = []
    current_top = None
    for w in words:
        if current_top is None or abs(w["top"] - current_top) <= 3:
            current.append(w)
            current_top = w["top"] if current_top is None else current_top
        else:
            lines.append(current)
            current = [w]
            current_top = w["top"]
    if current:
        lines.append(current)

    result = []
    for line_words in lines:
        line_words.sort(key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in line_words)
        x0 = min(w["x0"] for w in line_words)
        x1 = max(w["x1"] for w in line_words)
        result.append((text, x0, x1))
    return result


def find_centered_lines(lines, page_width):
    """Решает, какие строки страницы визуально центрированы.

    Если на странице есть настоящий "рабочий" левый край абзацев (он же —
    самое частое значение x0), сравниваем каждую строку с полями именно
    этого текстового столбца — так длинные строки списков (которые почти
    doезжают до правого края страницы) не путаются с реальным центром.
    Если страница почти целиком из заголовков без обычных абзацев (как
    титульный лист) — такого "рабочего края" нет, и надёжнее сравнивать
    с геометрическим центром страницы.
    """
    if not lines:
        return set()

    x0_counts = Counter(round(x0) for _, x0, _ in lines)
    body_left, freq = x0_counts.most_common(1)[0]
    use_body_margin = freq >= max(3, len(lines) * 0.3)

    centered = set()
    if use_body_margin:
        rights = [x1 for _, x0, x1 in lines if abs(x0 - body_left) <= 2]
        body_right = max(rights) if rights else max(x1 for _, _, x1 in lines)
        body_width = body_right - body_left
        for i, (_, x0, x1) in enumerate(lines):
            left_gap = x0 - body_left
            right_gap = body_right - x1
            line_width = x1 - x0
            if (
                left_gap > 8
                and right_gap > 8
                and abs(left_gap - right_gap) < max(body_width * 0.05, 10)
                and line_width < body_width * 0.9
            ):
                centered.add(i)
    else:
        for i, (_, x0, x1) in enumerate(lines):
            center = (x0 + x1) / 2
            line_width = x1 - x0
            if (
                abs(center - page_width / 2) < page_width * 0.08
                and line_width < page_width * 0.85
            ):
                centered.add(i)
    return centered


def pdf_to_docx(input_bytes):
    doc = Document()
    with pdfplumber.open(io.BytesIO(input_bytes)) as pdf:
        for i, page in enumerate(pdf.pages):
            lines = extract_lines_with_bbox(page)
            centered_idx = find_centered_lines(lines, page.width)
            for j, (text, _, _) in enumerate(lines):
                if not text.strip():
                    continue
                paragraph = doc.add_paragraph(text)
                if j in centered_idx:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < len(pdf.pages) - 1:
                doc.add_page_break()
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def docx_to_pdf(input_bytes):
    doc = Document(io.BytesIO(input_bytes))
    buffer = io.BytesIO()
    canvas = pdf_canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = 2 * cm
    y = height - margin
    line_height = 16
    font_size = 11
    max_chars = 90
    canvas.setFont("DejaVuSans", font_size)

    for para in doc.paragraphs:
        text = para.text
        lines = textwrap.wrap(text, max_chars) if text.strip() else [""]
        for line in lines:
            if y < margin:
                canvas.showPage()
                canvas.setFont("DejaVuSans", font_size)
                y = height - margin
            canvas.drawString(margin, y, line)
            y -= line_height
        y -= line_height * 0.4

    canvas.save()
    buffer.seek(0)
    return buffer

# ---------- Настоящий AI-апскейл через FSRCNN (лёгкая модель, CPU) ----------

MODELS_DIR = os.path.join(os.path.dirname(__file__), "models")
MODEL_URLS = {
    2: "https://raw.githubusercontent.com/Saafke/FSRCNN_Tensorflow/master/models/FSRCNN_x2.pb",
    3: "https://raw.githubusercontent.com/Saafke/FSRCNN_Tensorflow/master/models/FSRCNN_x3.pb",
    4: "https://raw.githubusercontent.com/Saafke/FSRCNN_Tensorflow/master/models/FSRCNN_x4.pb",
}

_sr_instance = cv2.dnn_superres.DnnSuperResImpl_create()
_loaded_scale = None


def ensure_model_downloaded(scale):
    os.makedirs(MODELS_DIR, exist_ok=True)
    path = os.path.join(MODELS_DIR, f"FSRCNN_x{scale}.pb")
    if not os.path.exists(path):
        print(f"Скачиваю модель FSRCNN x{scale}...")
        urllib.request.urlretrieve(MODEL_URLS[scale], path)
    return path


def get_sr_model(scale):
    global _loaded_scale
    if _loaded_scale != scale:
        path = ensure_model_downloaded(scale)
        _sr_instance.readModel(path)
        _sr_instance.setModel("fsrcnn", scale)
        _loaded_scale = scale
    return _sr_instance


def check_limit():
    ip = request.remote_addr
    now = time.time()
    entry = usage.get(ip)

    if not entry or now > entry["reset_at"]:
        entry = {"count": 0, "reset_at": now + DAY_SECONDS}

    if entry["count"] >= DAILY_LIMIT:
        minutes_left = int((entry["reset_at"] - now) / 60) + 1
        return jsonify({
            "error": "limit_reached",
            "message": f"Лимит бесплатных обработок исчерпан ({DAILY_LIMIT}/день). Попробуй через {minutes_left} мин.",
        }), 429

    entry["count"] += 1
    usage[ip] = entry
    return None


@app.route("/process", methods=["POST"])
def process():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "Файл не найден"}), 400

    limit_response = check_limit()
    if limit_response:
        return limit_response

    mode = request.form.get("mode")  # 'compress' | 'bg-remove' | 'upscale' | 'convert' | 'convert-doc'
    compression = request.form.get("compression", "medium")
    scale = int(request.form.get("scale", 2))
    target_format = request.form.get("target_format", "png").lower()
    tier = request.form.get("tier", "basic")

    input_bytes = file.read()

    if mode == "convert-doc":
        if tier != "basic":
            return jsonify({
                "error": "pro_unavailable",
                "message": "Улучшенная конвертация документов временно недоступна. Попробуй Упрощённую.",
            }), 400

        if target_format not in ("pdf", "docx"):
            return jsonify({"error": "unsupported_format", "message": "Неподдерживаемый формат"}), 400

        source_name = (file.filename or "").lower()
        try:
            if source_name.endswith(".pdf") and target_format == "docx":
                buffer = pdf_to_docx(input_bytes)
                mimetype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif source_name.endswith((".docx", ".doc")) and target_format == "pdf":
                buffer = docx_to_pdf(input_bytes)
                mimetype = "application/pdf"
            else:
                return jsonify({
                    "error": "unsupported_combo",
                    "message": "В Упрощённой доступно только PDF → Word и Word → PDF",
                }), 400
        except Exception:
            return jsonify({"error": "Не удалось прочитать документ. Убедись, что PDF содержит текст, а не скан."}), 400

        return send_file(buffer, mimetype=mimetype, as_attachment=True,
                          download_name=f"result.{target_format}")

    if mode == "convert":
        if target_format not in CONVERT_FORMATS:
            return jsonify({"error": "unsupported_format", "message": "Неподдерживаемый формат"}), 400

        pil_format, mimetype = CONVERT_FORMATS[target_format]
        try:
            image = Image.open(io.BytesIO(input_bytes))
        except Exception:
            return jsonify({"error": "Не удалось прочитать изображение"}), 400

        buffer = io.BytesIO()
        if pil_format == "JPEG":
            # JPEG не поддерживает прозрачность — подкладываем белый фон
            image = image.convert("RGBA")
            background = Image.new("RGB", image.size, (255, 255, 255))
            background.paste(image, mask=image.split()[3])
            background.save(buffer, format="JPEG", quality=92)
        elif pil_format == "ICO":
            image = image.convert("RGBA")
            image.save(buffer, format="ICO", sizes=[(16, 16), (32, 32), (48, 48), (64, 64)])
        else:
            image.save(buffer, format=pil_format)

        buffer.seek(0)
        return send_file(buffer, mimetype=mimetype)

    if mode == "bg-remove":
        result_bytes = remove(input_bytes)
        return send_file(io.BytesIO(result_bytes), mimetype="image/png")

    if mode == "compress":
        image = Image.open(io.BytesIO(input_bytes)).convert("RGB")
        quality = QUALITY_MAP.get(compression, 60)
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=quality, optimize=True)
        buffer.seek(0)
        return send_file(buffer, mimetype="image/jpeg")

    if mode == "upscale":
        if scale not in (2, 3, 4):
            scale = 2

        np_arr = np.frombuffer(input_bytes, np.uint8)
        cv_image = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)

        if cv_image is None:
            return jsonify({"error": "Не удалось прочитать изображение"}), 400

        h, w = cv_image.shape[:2]
        if max(h, w) > MAX_INPUT_SIDE:
            return jsonify({
                "error": "image_too_large",
                "message": "Фото слишком большое для увеличения. Максимальная сторона — 2000px.",
            }), 400

        sr = get_sr_model(scale)
        upscaled = sr.upsample(cv_image)

        success, buffer = cv2.imencode(".jpg", upscaled, [cv2.IMWRITE_JPEG_QUALITY, 92])
        if not success:
            return jsonify({"error": "Не удалось сохранить результат"}), 500

        return send_file(io.BytesIO(buffer.tobytes()), mimetype="image/jpeg")

    return jsonify({"error": "Неизвестный режим"}), 400


@app.route("/health")
def health():
    return jsonify({"ok": True})


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "false") == "true"
    port = int(os.environ.get("PORT", 5001))
    app.run(host="0.0.0.0", port=port, debug=debug)
